#!/usr/bin/env python3
"""Build ../plots.xlsx and ../plots.pdf from the benchmark logs in ../runs/.

Every benchmark log contains machine-readable lines
    CSV,alg,log2N,rank,extent,threads,rep,seconds,GBps
with GBps = 2*N bytes / seconds / 1e9 (N bytes loaded + N bytes stored).
Reported bandwidth = best repetition (minimum time), as STREAM does.

Usage:  python3 plots.py            (run from src/, or anywhere)
"""
import collections
import os
import re
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import Reference, ScatterChart, Series  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RUNS = os.path.join(ROOT, "runs")

# fixed categorical order (one slot per rank, never cycled) + a marker per
# series so identity never relies on color alone
COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300",
          "#4a3aa7", "#e34948"]
MARKERS = ["o", "s", "^", "D", "v", "P", "X", "*"]
XL_MARKERS = ["circle", "square", "triangle", "diamond", "x", "star", "plus", "dash"]


def read_log(name):
    """-> list of dicts, one per CSV line; `kernel` from '### BLOCK_KERNEL=k' headers."""
    path = os.path.join(RUNS, name)
    rows, kernel = [], None
    if not os.path.exists(path):
        return rows
    with open(path) as f:
        for line in f:
            m = re.match(r"### BLOCK_KERNEL=(\d+)", line)
            if m:
                kernel = int(m.group(1))
            if not line.startswith("CSV,"):
                continue
            _, alg, lg, rank, ext, thr, rep, sec, gbps = line.strip().split(",")
            rows.append(dict(alg=alg, log2N=int(lg), rank=int(rank), extent=int(ext),
                             threads=int(thr), rep=int(rep), seconds=float(sec),
                             gbps=float(gbps), kernel=kernel))
    return rows


def best(rows, *keys):
    """Group by keys -> (min seconds, max GB/s, median seconds)."""
    groups = collections.defaultdict(list)
    for r in rows:
        groups[tuple(r[k] for k in keys)].append(r)
    out = {}
    for k, g in groups.items():
        secs = sorted(r["seconds"] for r in g)
        out[k] = (secs[0], max(r["gbps"] for r in g), secs[len(secs) // 2])
    return out


def scaling_sheet(wb, title, rows, note):
    """Table cores x rank of GB/s + an Excel scatter-with-lines chart."""
    ws = wb.create_sheet(title)
    agg = best(rows, "rank", "threads")
    ranks = sorted({r for r, _ in agg})
    threads = sorted({t for _, t in agg})
    lg = rows[0]["log2N"]
    ws["A1"] = f"{title}: measured bandwidth [GB/s] = 2N / t_min, N = 2^{lg} bytes"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = note
    ws.cell(row=4, column=1, value="cores")
    for j, rk in enumerate(ranks):
        ext = next(r["extent"] for r in rows if r["rank"] == rk)
        ws.cell(row=4, column=2 + j, value=f"rank {rk} ({ext}^{rk})")
    for i, t in enumerate(threads):
        ws.cell(row=5 + i, column=1, value=t)
        for j, rk in enumerate(ranks):
            if (rk, t) in agg:
                ws.cell(row=5 + i, column=2 + j, value=round(agg[(rk, t)][1], 3))
    # runtime table next to it
    c0 = len(ranks) + 3
    ws.cell(row=4, column=c0, value="cores")
    for j, rk in enumerate(ranks):
        ws.cell(row=4, column=c0 + 1 + j, value=f"t_min [s] rank {rk}")
    for i, t in enumerate(threads):
        ws.cell(row=5 + i, column=c0, value=t)
        for j, rk in enumerate(ranks):
            if (rk, t) in agg:
                ws.cell(row=5 + i, column=c0 + 1 + j, value=round(agg[(rk, t)][0], 4))

    ch = ScatterChart()
    ch.title = f"{title}: bandwidth vs cores (one line per rank)"
    ch.style = 13
    ch.x_axis.title = "cores"
    ch.y_axis.title = "bandwidth [GB/s]"
    ch.height, ch.width = 10, 18
    xref = Reference(ws, min_col=1, min_row=5, max_row=4 + len(threads))
    for j in range(len(ranks)):
        yref = Reference(ws, min_col=2 + j, min_row=4, max_row=4 + len(threads))
        s = Series(yref, xref, title_from_data=True)
        s.graphicalProperties.line.solidFill = COLORS[j % 8][1:]
        s.graphicalProperties.line.width = 22000
        s.marker.symbol = XL_MARKERS[j % 8]
        s.marker.size = 7
        s.marker.graphicalProperties.solidFill = COLORS[j % 8][1:]
        s.marker.graphicalProperties.line.solidFill = COLORS[j % 8][1:]
        s.smooth = False
        ch.series.append(s)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, f"A{7 + len(threads)}")
    return agg, ranks, threads


def scaling_figure(pdf, title, agg, ranks, threads):
    fig, ax = plt.subplots(figsize=(8.5, 5.2))
    for j, rk in enumerate(ranks):
        xs = [t for t in threads if (rk, t) in agg]
        ys = [agg[(rk, t)][1] for t in xs]
        ax.plot(xs, ys, color=COLORS[j % 8], marker=MARKERS[j % 8], lw=2, ms=7,
                label=f"rank {rk}")
        ax.annotate(f"rank {rk}", (xs[-1], ys[-1]), xytext=(6, 0),
                    textcoords="offset points", va="center", fontsize=8, color="#52514e")
    ax.set_xlabel("cores (OpenMP threads, one per core)")
    ax.set_ylabel("bandwidth 2N/t  [GB/s]")
    ax.set_title(title, fontsize=11)
    ax.set_xticks(threads if len(threads) <= 16 else threads[:: max(1, len(threads) // 16)])
    ax.set_ylim(bottom=0)
    ax.grid(True, color="#e6e5e0", lw=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    fig.tight_layout()
    pdf.savefig(fig)
    plt.close(fig)


def main():
    machine = ""
    mpath = os.path.join(RUNS, "machine.txt")
    if os.path.exists(mpath):
        txt = open(mpath).read()
        m = re.search(r"Model name:\s*(.+)", txt)
        c = re.search(r"cores used: (\d+)", txt)
        machine = (m.group(1).strip() if m else "") + (f", {c.group(1)} cores" if c else "")

    wb = Workbook()
    wb.remove(wb.active)
    pdf = PdfPages(os.path.join(ROOT, "plots.pdf"))
    made = []

    for log, title, note in [
        ("bench_iterative.txt", "iterative",
         "Main measurement: iterative (odometer) reorder, parallelized by splitting the output."),
        ("bench_blocked.txt", "blocked (Part 3)",
         "Part 3: cache-blocked reorder (64x64 byte tiles, SSE2 transposes, streaming stores)."),
    ]:
        rows = read_log(log)
        if not rows:
            continue
        lg = rows[0]["log2N"]
        rows = [r for r in rows if r["log2N"] == lg]
        agg, ranks, threads = scaling_sheet(wb, title[:31], rows, note)
        scaling_figure(pdf, f"{title}: bandwidth vs cores, N = 2^{lg} bytes"
                       + (f"\n{machine}" if machine else ""), agg, ranks, threads)
        made.append(title)

    # all algorithms at the maximum core count
    rows = read_log("bench_all_algorithms.txt") + [
        r for f in ("bench_iterative.txt", "bench_blocked.txt") for r in read_log(f)]
    if rows:
        pmax = max(r["threads"] for r in read_log("bench_all_algorithms.txt") or rows)
        lg = max(r["log2N"] for r in rows)
        rows = [r for r in rows if r["threads"] == pmax and r["log2N"] == lg]
        agg = best(rows, "alg", "rank")
        algs = [a for a in ["index", "iterative", "recursive", "recursive_nd", "stages",
                            "blocked"] if any(k[0] == a for k in agg)]
        ranks = sorted({k[1] for k in agg})
        ws = wb.create_sheet("all algorithms")
        ws["A1"] = f"All algorithms at {pmax} cores, N = 2^{lg}: bandwidth 2N/t_min [GB/s] and t_min [s]"
        ws["A1"].font = Font(bold=True)
        ws.cell(row=3, column=1, value="algorithm")
        for j, rk in enumerate(ranks):
            ws.cell(row=3, column=2 + j, value=f"GB/s rank {rk}")
            ws.cell(row=3, column=3 + len(ranks) + j, value=f"t[s] rank {rk}")
        for i, a in enumerate(algs):
            ws.cell(row=4 + i, column=1, value=a)
            for j, rk in enumerate(ranks):
                if (a, rk) in agg:
                    ws.cell(row=4 + i, column=2 + j, value=round(agg[(a, rk)][1], 3))
                    ws.cell(row=4 + i, column=3 + len(ranks) + j, value=round(agg[(a, rk)][0], 3))
        fig, ax = plt.subplots(figsize=(8.5, 5.2))
        for j, a in enumerate(algs):
            xs = [rk for rk in ranks if (a, rk) in agg]
            ax.plot(xs, [agg[(a, rk)][1] for rk in xs], color=COLORS[j % 8],
                    marker=MARKERS[j % 8], lw=2, ms=7, label=a)
        ax.set_xscale("log", base=2)
        ax.set_yscale("log")
        ax.set_xticks(ranks, [str(r) for r in ranks])
        ax.set_xlabel("tensor rank n (square shape, extent 2^(%d/n))" % lg)
        ax.set_ylabel("bandwidth 2N/t  [GB/s]  (log scale)")
        ax.set_title(f"All implementations at {pmax} cores, N = 2^{lg} bytes", fontsize=11)
        ax.grid(True, which="both", color="#e6e5e0", lw=0.8)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(frameon=False, fontsize=8)
        fig.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)
        made.append("all algorithms")

    # runtime as a function of problem size
    rows = read_log("size_sweep.txt")
    if rows:
        agg = best(rows, "alg", "log2N", "rank")
        ws = wb.create_sheet("size sweep")
        ws["A1"] = "Runtime t_min [s] vs problem size N (every square rank of each N)"
        ws["A1"].font = Font(bold=True)
        hdr = ["algorithm", "log2 N", "N bytes", "rank", "extent", "t_min [s]", "GB/s"]
        for j, h in enumerate(hdr):
            ws.cell(row=3, column=1 + j, value=h)
        i = 4
        for (a, lg, rk), (tmin, gb, _) in sorted(agg.items()):
            ext = next(r["extent"] for r in rows if (r["alg"], r["log2N"], r["rank"]) == (a, lg, rk))
            for j, v in enumerate([a, lg, 2 ** lg, rk, ext, round(tmin, 6), round(gb, 3)]):
                ws.cell(row=i, column=1 + j, value=v)
            i += 1
        algs = sorted({k[0] for k in agg})
        fig, axs = plt.subplots(1, len(algs), figsize=(5.0 * len(algs), 4.6), squeeze=False)
        for ax, a in zip(axs[0], algs):
            for j, rk in enumerate([1, 2, 4]):
                pts = sorted((lg, v[0]) for (aa, lg, r), v in agg.items() if aa == a and r == rk)
                if pts:
                    ax.plot([2 ** p[0] for p in pts], [p[1] for p in pts], color=COLORS[j],
                            marker=MARKERS[j], lw=2, ms=7, label=f"rank {rk}")
            ax.set_xscale("log", base=2)
            ax.set_yscale("log")
            ax.set_xlabel("N [bytes]")
            ax.set_ylabel("runtime t_min [s]")
            ax.set_title(f"{a}: runtime vs problem size", fontsize=11)
            ax.grid(True, which="both", color="#e6e5e0", lw=0.8)
            ax.spines[["top", "right"]].set_visible(False)
            ax.legend(frameon=False, fontsize=8)
        fig.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)
        made.append("size sweep")

    # Part 3 kernel ladder
    rows = read_log("part3_kernel_ladder.txt")
    if rows:
        agg = best(rows, "kernel", "rank", "threads")
        kernels = sorted({k[0] for k in agg})
        ranks = sorted({k[1] for k in agg})
        tmax = max(k[2] for k in agg)
        ws = wb.create_sheet("Part 3 kernels")
        ws["A1"] = f"Part 3: tile kernel ladder, bandwidth [GB/s] at {tmax} cores"
        ws["A1"].font = Font(bold=True)
        names = {0: "0 scalar direct", 1: "1 + L1 buffer", 2: "2 + SSE2 16x16",
                 3: "3 + streaming stores", 4: "4 + wider input rows"}
        ws.cell(row=3, column=1, value="kernel")
        for j, rk in enumerate(ranks):
            ws.cell(row=3, column=2 + j, value=f"rank {rk}")
        for i, k in enumerate(kernels):
            ws.cell(row=4 + i, column=1, value=names.get(k, str(k)))
            for j, rk in enumerate(ranks):
                if (k, rk, tmax) in agg:
                    ws.cell(row=4 + i, column=2 + j, value=round(agg[(k, rk, tmax)][1], 3))
        fig, ax = plt.subplots(figsize=(8.5, 5.2))
        for j, k in enumerate(kernels):
            xs = [rk for rk in ranks if (k, rk, tmax) in agg]
            ax.plot(xs, [agg[(k, rk, tmax)][1] for rk in xs], color=COLORS[j],
                    marker=MARKERS[j], lw=2, ms=7, label=names.get(k, str(k)))
        ax.set_xscale("log", base=2)
        ax.set_xticks(ranks, [str(r) for r in ranks])
        ax.set_ylim(bottom=0)
        ax.set_xlabel("tensor rank n")
        ax.set_ylabel("bandwidth 2N/t  [GB/s]")
        ax.set_title(f"Part 3: blocking step by step ({tmax} cores)", fontsize=11)
        ax.grid(True, color="#e6e5e0", lw=0.8)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(frameon=False, fontsize=8)
        fig.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)
        made.append("Part 3 kernels")

    # raw data
    ws = wb.create_sheet("raw")
    hdr = ["file", "kernel", "alg", "log2N", "rank", "extent", "threads", "rep", "seconds", "GBps"]
    ws.append(hdr)
    for f in sorted(os.listdir(RUNS)) if os.path.isdir(RUNS) else []:
        if f.endswith(".txt"):
            for r in read_log(f):
                ws.append([f, r["kernel"], r["alg"], r["log2N"], r["rank"], r["extent"],
                           r["threads"], r["rep"], r["seconds"], r["gbps"]])

    pdf.close()
    if not made:
        print("no benchmark logs found in", RUNS, file=sys.stderr)
        return 1
    wb.save(os.path.join(ROOT, "plots.xlsx"))
    print("wrote plots.xlsx and plots.pdf:", ", ".join(made))
    return 0


if __name__ == "__main__":
    sys.exit(main())
